"""
Transcription Export Module

Handles exporting transcriptions to various formats (DOCX, SRT, VTT, JSON).
Extracted from app.workers.tasks.transcription for better modularity.
"""
import json
from typing import Dict, Any, Optional, List
from datetime import datetime, timedelta
from app.core.logging_config import get_logger

logger = get_logger(__name__)


class TranscriptionExporter:
    """Handles exporting transcriptions to various formats."""

    @staticmethod
    def export_to_docx(segments: List[Dict[str, Any]], output_path: str) -> bool:
        """
        Export transcription to DOCX format.

        Args:
            segments: List of transcription segments
            output_path: Path to output DOCX file

        Returns:
            True if successful, False otherwise
        """
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document()

            # Add title
            title = doc.add_heading('Transcription', 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Add metadata
            doc.add_paragraph(f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')}")
            doc.add_paragraph("")

            # Group segments by speaker
            current_speaker = None
            current_paragraph = None

            for segment in segments:
                speaker = segment.get('speaker', 'SPEAKER_01')
                start_time = segment.get('start', 0.0)
                text = segment.get('text', '').strip()

                if not text:
                    continue

                # Add speaker header if speaker changed
                if speaker != current_speaker:
                    current_speaker = speaker

                    # Add speaker heading
                    speaker_para = doc.add_paragraph()
                    speaker_run = speaker_para.add_run(f"\n{speaker}")
                    speaker_run.bold = True
                    speaker_run.font.size = Pt(12)
                    speaker_run.font.color.rgb = RGBColor(0, 0, 139)

                # Add timestamped text
                para = doc.add_paragraph()

                # Add timestamp
                timestamp_str = TranscriptionExporter._format_timestamp(start_time)
                time_run = para.add_run(f"[{timestamp_str}] ")
                time_run.font.color.rgb = RGBColor(128, 128, 128)
                time_run.font.size = Pt(9)

                # Add text
                text_run = para.add_run(text)
                text_run.font.size = Pt(11)

            doc.save(output_path)
            logger.info(f"Exported DOCX to {output_path}")
            return True

        except Exception as e:
            logger.error(f"Failed to export DOCX: {e}")
            return False

    @staticmethod
    def export_to_srt(segments: List[Dict[str, Any]], output_path: str) -> bool:
        """
        Export transcription to SRT subtitle format.

        Args:
            segments: List of transcription segments
            output_path: Path to output SRT file

        Returns:
            True if successful, False otherwise
        """
        try:
            with open(output_path, 'w', encoding='utf-8') as f:
                for idx, segment in enumerate(segments, start=1):
                    start = segment.get('start', 0.0)
                    end = segment.get('end', 0.0)
                    text = segment.get('text', '').strip()
                    speaker = segment.get('speaker', '')

                    if not text:
                        continue

                    # Format: sequence number
                    f.write(f"{idx}\n")

                    # Format: start --> end
                    start_str = TranscriptionExporter._format_srt_timestamp(start)
                    end_str = TranscriptionExporter._format_srt_timestamp(end)
                    f.write(f"{start_str} --> {end_str}\n")

                    # Format: text (with optional speaker)
                    if speaker:
                        f.write(f"[{speaker}] {text}\n")
                    else:
                        f.write(f"{text}\n")

                    # Blank line between subtitles
                    f.write("\n")

            logger.info(f"Exported SRT to {output_path}")
            return True

        except Exception as e:
            logger.error(f"Failed to export SRT: {e}")
            return False

    @staticmethod
    def export_to_vtt(segments: List[Dict[str, Any]], output_path: str) -> bool:
        """
        Export transcription to WebVTT subtitle format.

        Args:
            segments: List of transcription segments
            output_path: Path to output VTT file

        Returns:
            True if successful, False otherwise
        """
        try:
            with open(output_path, 'w', encoding='utf-8') as f:
                # WebVTT header
                f.write("WEBVTT\n\n")

                for idx, segment in enumerate(segments, start=1):
                    start = segment.get('start', 0.0)
                    end = segment.get('end', 0.0)
                    text = segment.get('text', '').strip()
                    speaker = segment.get('speaker', '')

                    if not text:
                        continue

                    # Format: start --> end
                    start_str = TranscriptionExporter._format_vtt_timestamp(start)
                    end_str = TranscriptionExporter._format_vtt_timestamp(end)
                    f.write(f"{start_str} --> {end_str}\n")

                    # Format: text with speaker as voice tag
                    if speaker:
                        f.write(f"<v {speaker}>{text}</v>\n")
                    else:
                        f.write(f"{text}\n")

                    # Blank line between cues
                    f.write("\n")

            logger.info(f"Exported VTT to {output_path}")
            return True

        except Exception as e:
            logger.error(f"Failed to export VTT: {e}")
            return False

    @staticmethod
    def export_to_json(
        segments: List[Dict[str, Any]],
        output_path: str,
        metadata: Optional[Dict[str, Any]] = None
    ) -> bool:
        """
        Export transcription to JSON format.

        Args:
            segments: List of transcription segments
            output_path: Path to output JSON file
            metadata: Optional metadata to include

        Returns:
            True if successful, False otherwise
        """
        try:
            data = {
                'metadata': metadata or {},
                'segments': segments,
                'generated_at': datetime.utcnow().isoformat()
            }

            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(data, f, indent=2, ensure_ascii=False)

            logger.info(f"Exported JSON to {output_path}")
            return True

        except Exception as e:
            logger.error(f"Failed to export JSON: {e}")
            return False

    @staticmethod
    def _format_timestamp(seconds: float) -> str:
        """Format seconds to HH:MM:SS."""
        td = timedelta(seconds=seconds)
        hours = int(td.total_seconds() // 3600)
        minutes = int((td.total_seconds() % 3600) // 60)
        secs = int(td.total_seconds() % 60)
        return f"{hours:02d}:{minutes:02d}:{secs:02d}"

    @staticmethod
    def _format_srt_timestamp(seconds: float) -> str:
        """Format seconds to SRT timestamp format (HH:MM:SS,mmm)."""
        td = timedelta(seconds=seconds)
        hours = int(td.total_seconds() // 3600)
        minutes = int((td.total_seconds() % 3600) // 60)
        secs = int(td.total_seconds() % 60)
        millis = int((seconds % 1) * 1000)
        return f"{hours:02d}:{minutes:02d}:{secs:02d},{millis:03d}"

    @staticmethod
    def _format_vtt_timestamp(seconds: float) -> str:
        """Format seconds to WebVTT timestamp format (HH:MM:SS.mmm)."""
        td = timedelta(seconds=seconds)
        hours = int(td.total_seconds() // 3600)
        minutes = int((td.total_seconds() % 3600) // 60)
        secs = int(td.total_seconds() % 60)
        millis = int((seconds % 1) * 1000)
        return f"{hours:02d}:{minutes:02d}:{secs:02d}.{millis:03d}"
