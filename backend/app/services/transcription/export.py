"""Export functionality for transcription service - all export formats."""

import io
import json
from sqlalchemy.orm import Session
from docx import Document as DocxDocument

from app.models.transcription import Transcription

logger = get_logger(__name__)


class TranscriptionExportService:
    """Export transcriptions in multiple formats."""

    @staticmethod
    def export_as_json(transcription: Transcription, db: Session) -> bytes:
        """Export transcription as JSON."""
        data = {
            "transcription_id": transcription.id,
            "case_id": transcription.case_id,
            "filename": transcription.filename,
            "format": transcription.format,
            "duration": transcription.duration,
            "speakers": transcription.speakers,
            "segments": transcription.segments,
            "created_at": transcription.created_at.isoformat(),
            "uploaded_at": transcription.uploaded_at.isoformat() if transcription.uploaded_at else None,
        }

        return json.dumps(data, indent=2).encode("utf-8")

    @staticmethod
    def export_as_docx(transcription: Transcription, db: Session) -> bytes:
        """Export transcription as DOCX with formatting."""
        doc = DocxDocument()

        # Add title
        title = doc.add_heading(
            f"Transcription: {transcription.filename}", 0
        )

        # Add metadata
        doc.add_paragraph(f"Transcription ID: {transcription.id}")
        doc.add_paragraph(f"Case ID: {transcription.case_id}")
        doc.add_paragraph(f"Format: {transcription.format or 'Unknown'}")
        if transcription.duration:
            minutes = int(transcription.duration // 60)
            seconds = int(transcription.duration % 60)
            doc.add_paragraph(f"Duration: {minutes}m {seconds}s")
        doc.add_paragraph(f"Created: {transcription.created_at.strftime('%Y-%m-%d %H:%M:%S')}")
        if transcription.uploaded_at:
            doc.add_paragraph(f"Uploaded: {transcription.uploaded_at.strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph()

        # Add speakers section if available
        if transcription.speakers:
            doc.add_heading("Speakers", level=1)
            for speaker in transcription.speakers:
                speaker_text = f"• {speaker.get('speaker_id', 'Unknown')}"
                if speaker.get("label"):
                    speaker_text += f" - {speaker['label']}"
                doc.add_paragraph(speaker_text)
            doc.add_paragraph()

        # Add transcription content
        doc.add_heading("Transcription", level=1)

        segments = transcription.segments or []
        for segment in segments:
            # Format timestamp
            start_time = TranscriptionExportService._format_timestamp(segment.get("start", 0))
            end_time = TranscriptionExportService._format_timestamp(segment.get("end", 0))
            timestamp = f"[{start_time} - {end_time}]"

            # Create paragraph with speaker and timestamp
            p = doc.add_paragraph()
            if segment.get("speaker"):
                p.add_run(f"{segment['speaker']}: ").bold = True
            p.add_run(f"{timestamp} ")
            p.add_run(segment.get("text", ""))
            p.add_run("\n")

        # Save to bytes
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.read()

    @staticmethod
    def export_as_srt(transcription: Transcription) -> bytes:
        """Export transcription as SRT subtitle format."""
        srt_content = []
        segments = transcription.segments or []

        for idx, segment in enumerate(segments, 1):
            start = TranscriptionExportService._format_srt_timestamp(segment.get("start", 0))
            end = TranscriptionExportService._format_srt_timestamp(segment.get("end", 0))
            text = segment.get("text", "")

            # Add speaker label if available
            if segment.get("speaker"):
                text = f"[{segment['speaker']}] {text}"

            srt_content.append(f"{idx}")
            srt_content.append(f"{start} --> {end}")
            srt_content.append(text)
            srt_content.append("")  # Empty line between subtitles

        return "\n".join(srt_content).encode("utf-8")

    @staticmethod
    def export_as_vtt(transcription: Transcription) -> bytes:
        """Export transcription as VTT subtitle format."""
        vtt_content = ["WEBVTT", ""]
        segments = transcription.segments or []

        for segment in segments:
            start = TranscriptionExportService._format_vtt_timestamp(segment.get("start", 0))
            end = TranscriptionExportService._format_vtt_timestamp(segment.get("end", 0))
            text = segment.get("text", "")

            # Add speaker label if available
            if segment.get("speaker"):
                text = f"<v {segment['speaker']}>{text}"

            vtt_content.append(f"{start} --> {end}")
            vtt_content.append(text)
            vtt_content.append("")  # Empty line between subtitles

        return "\n".join(vtt_content).encode("utf-8")

    @staticmethod
    def export_as_txt(transcription: Transcription) -> bytes:
        """Export transcription as plain text."""
        txt_content = []
        segments = transcription.segments or []

        for segment in segments:
            start_time = TranscriptionExportService._format_timestamp(segment.get("start", 0))
            text = segment.get("text", "")

            # Add speaker and timestamp
            line = f"[{start_time}]"
            if segment.get("speaker"):
                line += f" {segment['speaker']}:"
            line += f" {text}"

            txt_content.append(line)

        return "\n".join(txt_content).encode("utf-8")

    @staticmethod
    def _format_timestamp(seconds: float) -> str:
        """Format seconds as HH:MM:SS."""
        hours = int(seconds // 3600)
        minutes = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        return f"{hours:02d}:{minutes:02d}:{secs:02d}"

    @staticmethod
    def _format_srt_timestamp(seconds: float) -> str:
        """Format seconds as SRT timestamp (HH:MM:SS,mmm)."""
        hours = int(seconds // 3600)
        minutes = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        millis = int((seconds % 1) * 1000)
        return f"{hours:02d}:{minutes:02d}:{secs:02d},{millis:03d}"

    @staticmethod
    def _format_vtt_timestamp(seconds: float) -> str:
        """Format seconds as VTT timestamp (HH:MM:SS.mmm)."""
        hours = int(seconds // 3600)
        minutes = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        millis = int((seconds % 1) * 1000)
        return f"{hours:02d}:{minutes:02d}:{secs:02d}.{millis:03d}"
